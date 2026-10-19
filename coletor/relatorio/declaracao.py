"""Declaração editável do operador (.docx).

Fica fora do manifesto por ser editável, mas cita o hash dele, o que a
vincula àquele pacote.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt


def gerar(pasta: Path, dados: dict) -> Path:
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(11)

    doc.add_heading("Declaração de coleta de evidências digitais", level=1)
    qualificacao = f", {dados['qualificacao']}" if dados.get("qualificacao") else ""
    referencia = f", referente a {dados['referencia']}" if dados.get("referencia") else ""
    doc.add_paragraph(
        f"Eu, {dados['operador']}{qualificacao}, declaro que realizei a coleta identificada como "
        f"{dados['identificador']}{referencia}, entre {dados['inicio']} e {dados['fim']}, com o uso do "
        f"Fácil: Coletor de Evidências Digitais Intuitivo, versão {dados['ferramenta']['versao']}.")
    doc.add_paragraph("Foram coletados os seguintes itens:")
    for item in dados["itens"]:
        doc.add_paragraph(f"{item['numero']}: {item['nome_tipo']}, {item['descricao']}", style="List Bullet")
    doc.add_paragraph(
        "O conjunto de arquivos da coleta está descrito no manifesto cujo hash SHA-256 é "
        f"{dados['hash_manifesto']}, conforme o relatório que acompanha esta declaração.")
    doc.add_paragraph("\n\nLocal e data: ______________________________")
    doc.add_paragraph("\n\n______________________________\n" + dados["operador"])

    caminho = pasta / "Declaracao.docx"
    doc.save(caminho)
    return caminho
